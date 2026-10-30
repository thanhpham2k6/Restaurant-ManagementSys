# -*- coding: utf-8 -*-
"""
Chương trình xuất file Word (.docx) Báo cáo Bài tập lớn Hệ thống Quản lý Nhà hàng POS
Quy chuẩn trình bày: Trường Đại học Công nghệ Giao thông Vận tải (UTT)
Khoa Công nghệ Thông tin
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Thiết lập màu nền cho ô trong bảng"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))

def set_cell_margins(cell, top=140, bottom=140, left=160, right=160):
    """Thiết lập padding lề cho ô trong bảng (20 dxa = 1 pt)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Thiết lập đường kẻ viền thanh lịch cho toàn bộ bảng"""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_table_header(table, col_widths, headers, bg_color="1E3A8A"):
    """Định dạng hàng tiêu đề của bảng"""
    hdr_row = table.rows[0]
    tblPr = table._tbl.tblPr
    # Lặp lại header nếu bảng ngắt trang
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    for i, title in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.width = Inches(col_widths[i])
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=160, bottom=160, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(title)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

def format_data_rows(table, col_widths, rows_data, aligns=None):
    """Điền dữ liệu các hàng trong bảng với màu so le (zebra striping)"""
    for r_idx, r_data in enumerate(rows_data):
        row = table.add_row()
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            cell.width = Inches(col_widths[c_idx])
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.2
            
            align = WD_ALIGN_PARAGRAPH.LEFT
            if aligns and c_idx < len(aligns):
                align = aligns[c_idx]
            p.alignment = align
            
            run = p.add_run(str(val))
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(30, 41, 59)

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(8)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 58, 138) # Deep Blue
    return h

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 118, 110) # Teal
    return h

def add_heading_3(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(51, 65, 85) # Slate
    return h

def add_p(doc, text="", bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=5, line_spacing=1.35, indent=0.5):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if indent and align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        p.paragraph_format.first_line_indent = Inches(indent)
    
    if text:
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_bullet(doc, text, bold_prefix="", level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = "Times New Roman"
        r_bold.font.size = Pt(13)
        r_bold.font.bold = True
        r_bold.font.color.rgb = RGBColor(15, 23, 42)
        
    r_text = p.add_run(text)
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(13)
    r_text.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_callout(doc, title, content):
    """Hộp ghi chú nổi bật với dải viền bên trái"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="0F766E"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.25
    r1 = p.add_run(f"★ {title}\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(15, 118, 110)
    
    r2 = p.add_run(content)
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_figure_placeholder(doc, fig_number, caption, note="Ảnh chụp màn hình thực tế từ phần mềm"):
    """Tạo khung vị trí chèn hình ảnh giao diện / biểu đồ"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=180, bottom=180, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:left w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:bottom w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:right w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(f"[KHUNG ẢNH: {caption.upper()}]\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(71, 85, 105)
    
    r2 = p.add_run(f"({note})")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(10)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(148, 163, 184)
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run(f"Hình {fig_number}: {caption}")
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(11)
    r_cap.font.bold = True
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(30, 41, 59)

print("Modules and builders loaded.")
